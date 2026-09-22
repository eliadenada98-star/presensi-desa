from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    jsonify,
    send_file
)

from werkzeug.security import (
    generate_password_hash,
    check_password_hash
)

from database import (
    get_connection,
    init_database,
    execute,
    fetch_one,
    fetch_all
)

from datetime import datetime
from pathlib import Path
from io import BytesIO

import base64
import uuid

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


# =========================================================
# KONFIGURASI APLIKASI
# =========================================================

app = Flask(__name__)

app.secret_key = "presensi-desa-secret-key-2026"

BASE_DIR = Path(__file__).resolve().parent


# =========================================================
# FOLDER UPLOAD
# =========================================================

UPLOAD_FOLDER = (
    BASE_DIR
    / "static"
    / "uploads"
    / "presensi"
)

UPLOAD_FOLDER.mkdir(
    parents=True,
    exist_ok=True
)


# =========================================================
# INISIALISASI DATABASE
# =========================================================

init_database()


# =========================================================
# FUNGSI WAKTU
# =========================================================

def waktu_sekarang():
    return datetime.now()


def tanggal_sekarang():
    return datetime.now().strftime("%Y-%m-%d")


def jam_sekarang():
    return datetime.now().strftime("%H:%M:%S")


# =========================================================
# HALAMAN UTAMA
# =========================================================

@app.route("/")
def index():

    if "user_id" in session:

        if session.get("role") == "admin":
            return redirect(
                url_for("admin_dashboard")
            )

        if session.get("role") == "perangkat":
            return redirect(
                url_for("perangkat_dashboard")
            )

    return redirect(
        url_for("login")
    )


# =========================================================
# LOGIN
# =========================================================

@app.route(
    "/login",
    methods=["GET", "POST"]
)
def login():

    if request.method == "POST":

        username = (
            request.form
            .get("username", "")
            .strip()
        )

        password = (
            request.form
            .get("password", "")
        )

        if not username or not password:

            return render_template(
                "login.html",
                error=(
                    "Username dan password "
                    "wajib diisi."
                )
            )

        user = fetch_one("""
            SELECT *
            FROM users
            WHERE username = ?
        """, (
            username,
        ))

        if user is None:

            return render_template(
                "login.html",
                error=(
                    "Username atau password salah."
                )
            )

        if not check_password_hash(
            user["password"],
            password
        ):

            return render_template(
                "login.html",
                error=(
                    "Username atau password salah."
                )
            )

        session.clear()

        session["user_id"] = user["id"]
        session["username"] = user["username"]
        session["nama"] = user["nama"]
        session["role"] = user["role"]
        session["perangkat_id"] = user["perangkat_id"]

        if user["role"] == "admin":

            return redirect(
                url_for("admin_dashboard")
            )

        return redirect(
            url_for("perangkat_dashboard")
        )

    return render_template(
        "login.html"
    )


# =========================================================
# LOGOUT
# =========================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect(
        url_for("login")
    )


# =========================================================
# CEK ADMIN
# =========================================================

def admin_required():

    if "user_id" not in session:
        return False

    if session.get("role") != "admin":
        return False

    return True


# =========================================================
# DASHBOARD ADMIN
# =========================================================

@app.route("/admin/dashboard")
def admin_dashboard():

    if not admin_required():

        return redirect(
            url_for("login")
        )

    jumlah_perangkat = fetch_one("""
        SELECT COUNT(*) AS total
        FROM perangkat
    """)["total"]

    jumlah_presensi = fetch_one("""
        SELECT COUNT(*) AS total
        FROM presensi
    """)["total"]

    hadir_hari_ini = fetch_one("""
        SELECT COUNT(*) AS total
        FROM presensi
        WHERE tanggal = ?
    """, (
        tanggal_sekarang(),
    ))["total"]

    perangkat_aktif = fetch_one("""
        SELECT COUNT(*) AS total
        FROM perangkat
        WHERE status = 'Aktif'
    """)["total"]

    return render_template(
        "admin_dashboard.html",
        jumlah_perangkat=jumlah_perangkat,
        jumlah_presensi=jumlah_presensi,
        hadir_hari_ini=hadir_hari_ini,
        perangkat_aktif=perangkat_aktif
    )


# =========================================================
# MASTER PERANGKAT
# =========================================================

@app.route("/admin/perangkat")
def admin_perangkat():

    if not admin_required():

        return redirect(
            url_for("login")
        )

    perangkat = fetch_all("""
        SELECT
            perangkat.*,
            users.username

        FROM perangkat

        LEFT JOIN users
            ON users.perangkat_id =
               perangkat.id

        ORDER BY perangkat.id DESC
    """)

    return render_template(
        "master_perangkat.html",
        perangkat=perangkat
    )


# =========================================================
# TAMBAH PERANGKAT
# =========================================================

@app.route(
    "/admin/perangkat/tambah",
    methods=["GET", "POST"]
)
def tambah_perangkat():

    if not admin_required():

        return redirect(
            url_for("login")
        )

    if request.method == "POST":

        nama = (
            request.form
            .get("nama", "")
            .strip()
        )

        nik = (
            request.form
            .get("nik", "")
            .strip()
        )

        jabatan = (
            request.form
            .get("jabatan", "")
            .strip()
        )

        username = (
            request.form
            .get("username", "")
            .strip()
        )

        password = (
            request.form
            .get("password", "")
        )

        if not nama or not username or not password:

            return render_template(
                "tambah_perangkat.html",
                error=(
                    "Nama, username dan "
                    "password wajib diisi."
                )
            )

        if len(password) < 6:

            return render_template(
                "tambah_perangkat.html",
                error=(
                    "Password minimal 6 karakter."
                )
            )

        existing = fetch_one("""
            SELECT id
            FROM users
            WHERE username = ?
        """, (
            username,
        ))

        if existing:

            return render_template(
                "tambah_perangkat.html",
                error=(
                    "Username sudah digunakan."
                )
            )

        perangkat_id = execute("""
            INSERT INTO perangkat (
                nama,
                nik,
                jabatan,
                status
            )

            VALUES (
                ?,
                ?,
                ?,
                'Aktif'
            )
        """, (
            nama,
            nik,
            jabatan
        ))

        password_hash = (
            generate_password_hash(
                password
            )
        )

        execute("""
            INSERT INTO users (
                username,
                password,
                nama,
                role,
                perangkat_id
            )

            VALUES (
                ?,
                ?,
                ?,
                'perangkat',
                ?
            )
        """, (
            username,
            password_hash,
            nama,
            perangkat_id
        ))

        return redirect(
            url_for("admin_perangkat")
        )

    return render_template(
        "tambah_perangkat.html"
    )


# =========================================================
# ADMIN PRESENSI
# =========================================================

@app.route("/admin/presensi")
def admin_presensi():

    if not admin_required():

        return redirect(
            url_for("login")
        )

    data = fetch_all("""
        SELECT
            presensi.*,
            perangkat.nama,
            perangkat.jabatan

        FROM presensi

        LEFT JOIN perangkat
            ON perangkat.id =
               presensi.perangkat_id

        ORDER BY
            presensi.tanggal DESC,
            presensi.id DESC
    """)

    return render_template(
        "admin_presensi.html",
        data=data
    )


# =========================================================
# HAPUS PRESENSI
# =========================================================

@app.route(
    "/admin/presensi/hapus/<int:presensi_id>",
    methods=["POST"]
)
def hapus_presensi(presensi_id):

    if not admin_required():

        return redirect(
            url_for("login")
        )

    presensi = fetch_one("""
        SELECT *
        FROM presensi
        WHERE id = ?
    """, (
        presensi_id,
    ))

    if presensi is None:

        return redirect(
            url_for("admin_presensi")
        )

    foto_masuk = presensi["foto_masuk"]

    if foto_masuk:

        try:

            file_masuk = (
                BASE_DIR
                / "static"
                / foto_masuk
            )

            if file_masuk.exists():

                file_masuk.unlink()

        except Exception:

            pass

    foto_pulang = presensi["foto_pulang"]

    if foto_pulang:

        try:

            file_pulang = (
                BASE_DIR
                / "static"
                / foto_pulang
            )

            if file_pulang.exists():

                file_pulang.unlink()

        except Exception:

            pass

    execute("""
        DELETE FROM presensi
        WHERE id = ?
    """, (
        presensi_id,
    ))

    return redirect(
        url_for("admin_presensi")
    )


# =========================================================
# FUNGSI QUERY LAPORAN
# =========================================================

def ambil_data_laporan(dari="", sampai=""):

    query = """
        SELECT
            presensi.*,
            perangkat.nama,
            perangkat.nik,
            perangkat.jabatan

        FROM presensi

        LEFT JOIN perangkat
            ON perangkat.id =
               presensi.perangkat_id

        WHERE 1 = 1
    """

    params = []

    if dari:

        query += """
            AND presensi.tanggal >= ?
        """

        params.append(dari)

    if sampai:

        query += """
            AND presensi.tanggal <= ?
        """

        params.append(sampai)

    query += """
        ORDER BY
            presensi.tanggal DESC,
            presensi.id DESC
    """

    return fetch_all(
        query,
        tuple(params)
    )


# =========================================================
# ADMIN LAPORAN
# =========================================================

@app.route("/admin/laporan")
def admin_laporan():

    if not admin_required():

        return redirect(
            url_for("login")
        )

    dari = (
        request.args
        .get("dari", "")
        .strip()
    )

    sampai = (
        request.args
        .get("sampai", "")
        .strip()
    )

    laporan = ambil_data_laporan(
        dari,
        sampai
    )

    return render_template(
        "admin_laporan.html",
        laporan=laporan,
        dari=dari,
        sampai=sampai
    )


# =========================================================
# DOWNLOAD LAPORAN EXCEL
# =========================================================

@app.route("/admin/laporan/excel")
def download_laporan_excel():

    if not admin_required():

        return redirect(
            url_for("login")
        )

    dari = (
        request.args
        .get("dari", "")
        .strip()
    )

    sampai = (
        request.args
        .get("sampai", "")
        .strip()
    )

    laporan = ambil_data_laporan(
        dari,
        sampai
    )

    # ---------------------------------------------
    # BUAT WORKBOOK
    # ---------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Laporan Presensi"

    # ---------------------------------------------
    # JUDUL
    # ---------------------------------------------

    worksheet["A1"] = (
        "LAPORAN PRESENSI PERANGKAT DESA"
    )

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    worksheet.merge_cells(
        "A1:H1"
    )

    worksheet.row_dimensions[1].height = 28

    # ---------------------------------------------
    # PERIODE
    # ---------------------------------------------

    if dari and sampai:

        periode = (
            f"{dari} s/d {sampai}"
        )

    elif dari:

        periode = (
            f"Mulai {dari}"
        )

    elif sampai:

        periode = (
            f"Sampai {sampai}"
        )

    else:

        periode = "Semua Data"

    worksheet["A2"] = "Periode"

    worksheet["A2"].font = Font(
        bold=True
    )

    worksheet["B2"] = periode

    worksheet.merge_cells(
        "B2:H2"
    )

    # ---------------------------------------------
    # HEADER TABEL
    # ---------------------------------------------

    headers = [
        "No",
        "Tanggal",
        "Nama",
        "NIK",
        "Jabatan",
        "Jam Masuk",
        "Jam Pulang",
        "Status"
    ]

    header_row = 4

    for col, header in enumerate(
        headers,
        start=1
    ):

        cell = worksheet.cell(
            row=header_row,
            column=col,
            value=header
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="2563EB"
        )

    # ---------------------------------------------
    # ISI DATA
    # ---------------------------------------------

    row = 5

    for nomor, item in enumerate(
        laporan,
        start=1
    ):

        values = [
            nomor,
            item["tanggal"] or "-",
            item["nama"] or "-",
            item["nik"] or "-",
            item["jabatan"] or "-",
            item["jam_masuk"] or "-",
            item["jam_pulang"] or "-",
            item["status"] or "-"
        ]

        for col, value in enumerate(
            values,
            start=1
        ):

            cell = worksheet.cell(
                row=row,
                column=col,
                value=value
            )

            cell.alignment = Alignment(
                vertical="center"
            )

        row += 1

    # ---------------------------------------------
    # LEBAR KOLOM
    # ---------------------------------------------

    widths = {
        "A": 8,
        "B": 15,
        "C": 28,
        "D": 20,
        "E": 24,
        "F": 15,
        "G": 15,
        "H": 15
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    # ---------------------------------------------
    # FREEZE HEADER
    # ---------------------------------------------

    worksheet.freeze_panes = "A5"

    # ---------------------------------------------
    # AUTO FILTER
    # ---------------------------------------------

    if laporan:

        last_row = 4 + len(laporan)

        worksheet.auto_filter.ref = (
            f"A4:H{last_row}"
        )

    # ---------------------------------------------
    # SIMPAN KE MEMORY
    # ---------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=(
            "laporan_presensi.xlsx"
        ),
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# =========================================================
# DOWNLOAD LAPORAN WORD
# =========================================================

@app.route("/admin/laporan/word")
def download_laporan_word():

    if not admin_required():

        return redirect(
            url_for("login")
        )

    dari = (
        request.args
        .get("dari", "")
        .strip()
    )

    sampai = (
        request.args
        .get("sampai", "")
        .strip()
    )

    laporan = ambil_data_laporan(
        dari,
        sampai
    )

    # ---------------------------------------------
    # BUAT DOKUMEN WORD
    # ---------------------------------------------

    document = Document()

    # ---------------------------------------------
    # JUDUL
    # ---------------------------------------------

    title = document.add_heading(
        "LAPORAN PRESENSI PERANGKAT DESA",
        level=1
    )

    title.alignment = 1

    # ---------------------------------------------
    # PERIODE
    # ---------------------------------------------

    if dari and sampai:

        periode = (
            f"{dari} s/d {sampai}"
        )

    elif dari:

        periode = (
            f"Mulai {dari}"
        )

    elif sampai:

        periode = (
            f"Sampai {sampai}"
        )

    else:

        periode = "Semua Data"

    paragraph = document.add_paragraph()

    paragraph.alignment = 1

    run = paragraph.add_run(
        f"Periode: {periode}"
    )

    run.bold = True

    document.add_paragraph()

    # ---------------------------------------------
    # TABEL
    # ---------------------------------------------

    table = document.add_table(
        rows=1,
        cols=8
    )

    table.style = "Table Grid"

    headers = [
        "No",
        "Tanggal",
        "Nama",
        "NIK",
        "Jabatan",
        "Jam Masuk",
        "Jam Pulang",
        "Status"
    ]

    # ---------------------------------------------
    # HEADER WORD
    # ---------------------------------------------

    header_cells = table.rows[0].cells

    for index, header in enumerate(
        headers
    ):

        header_cells[index].text = header

        for paragraph in (
            header_cells[index].paragraphs
        ):

            for run in paragraph.runs:

                run.bold = True

                run.font.size = (
                    document.styles["Normal"]
                    .font.size
                )

    # ---------------------------------------------
    # DATA
    # ---------------------------------------------

    for nomor, item in enumerate(
        laporan,
        start=1
    ):

        cells = table.add_row().cells

        cells[0].text = str(
            nomor
        )

        cells[1].text = str(
            item["tanggal"] or "-"
        )

        cells[2].text = str(
            item["nama"] or "-"
        )

        cells[3].text = str(
            item["nik"] or "-"
        )

        cells[4].text = str(
            item["jabatan"] or "-"
        )

        cells[5].text = str(
            item["jam_masuk"] or "-"
        )

        cells[6].text = str(
            item["jam_pulang"] or "-"
        )

        cells[7].text = str(
            item["status"] or "-"
        )

    # ---------------------------------------------
    # TOTAL DATA
    # ---------------------------------------------

    document.add_paragraph()

    total_paragraph = document.add_paragraph()

    total_run = total_paragraph.add_run(
        f"Total data presensi: {len(laporan)}"
    )

    total_run.bold = True

    # ---------------------------------------------
    # SIMPAN KE MEMORY
    # ---------------------------------------------

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=(
            "laporan_presensi.docx"
        ),
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
    )


# =========================================================
# ADMIN PENGATURAN
# =========================================================

@app.route("/admin/pengaturan")
def admin_pengaturan():

    if not admin_required():

        return redirect(
            url_for("login")
        )

    return render_template(
        "admin_pengaturan.html"
    )


# =========================================================
# DASHBOARD PERANGKAT
# =========================================================

@app.route("/perangkat/dashboard")
def perangkat_dashboard():

    if "user_id" not in session:

        return redirect(
            url_for("login")
        )

    if session.get("role") != "perangkat":

        return redirect(
            url_for("admin_dashboard")
        )

    perangkat_id = session.get(
        "perangkat_id"
    )

    perangkat = fetch_one("""
        SELECT *
        FROM perangkat
        WHERE id = ?
    """, (
        perangkat_id,
    ))

    presensi_hari_ini = fetch_one("""
        SELECT *
        FROM presensi
        WHERE perangkat_id = ?
        AND tanggal = ?
        ORDER BY id DESC
        LIMIT 1
    """, (
        perangkat_id,
        tanggal_sekarang()
    ))

    return render_template(
        "perangkat_dashboard.html",
        perangkat=perangkat,
        presensi=presensi_hari_ini
    )


# =========================================================
# HALAMAN ABSEN
# =========================================================

@app.route("/perangkat/absen")
def perangkat_absen():

    if "user_id" not in session:

        return redirect(
            url_for("login")
        )

    if session.get("role") != "perangkat":

        return redirect(
            url_for("admin_dashboard")
        )

    perangkat_id = session.get(
        "perangkat_id"
    )

    tanggal = tanggal_sekarang()

    presensi = fetch_one("""
        SELECT *
        FROM presensi
        WHERE perangkat_id = ?
        AND tanggal = ?
        ORDER BY id DESC
        LIMIT 1
    """, (
        perangkat_id,
        tanggal
    ))

    print()
    print("========================================")
    print("DEBUG HALAMAN ABSEN")
    print("========================================")
    print("Perangkat ID :", perangkat_id)
    print("Tanggal      :", tanggal)

    if presensi:

        print(
            "Presensi ID  :",
            presensi["id"]
        )

        print(
            "Jam Masuk    :",
            presensi["jam_masuk"]
        )

        print(
            "Jam Pulang   :",
            presensi["jam_pulang"]
        )

    else:

        print(
            "Presensi      : BELUM ADA"
        )

    print("========================================")
    print()

    return render_template(
        "perangkat_absen.html",
        presensi=presensi
    )


# =========================================================
# SIMPAN ABSEN
# =========================================================

@app.route(
    "/perangkat/absen/simpan",
    methods=["POST"]
)
def simpan_absen():

    if "user_id" not in session:

        return jsonify({
            "success": False,
            "message": (
                "Silakan login terlebih dahulu."
            )
        }), 401

    if session.get("role") != "perangkat":

        return jsonify({
            "success": False,
            "message": "Akses ditolak."
        }), 403

    perangkat_id = session.get(
        "perangkat_id"
    )

    jenis = request.form.get(
        "jenis"
    )

    foto = request.form.get(
        "foto"
    )

    latitude = request.form.get(
        "latitude"
    )

    longitude = request.form.get(
        "longitude"
    )

    if jenis not in (
        "masuk",
        "pulang"
    ):

        return jsonify({
            "success": False,
            "message": (
                "Jenis presensi tidak valid."
            )
        })

    if not foto:

        return jsonify({
            "success": False,
            "message": (
                "Foto tidak ditemukan."
            )
        })

    tanggal = tanggal_sekarang()

    presensi = fetch_one("""
        SELECT *
        FROM presensi
        WHERE perangkat_id = ?
        AND tanggal = ?
        ORDER BY id DESC
        LIMIT 1
    """, (
        perangkat_id,
        tanggal
    ))

    # ---------------------------------------------
    # CEK ABSEN MASUK
    # ---------------------------------------------

    if jenis == "masuk":

        if (
            presensi
            and presensi["jam_masuk"]
        ):

            return jsonify({
                "success": False,
                "message": (
                    "Anda sudah melakukan "
                    "absen masuk hari ini."
                )
            })

    # ---------------------------------------------
    # CEK ABSEN PULANG
    # ---------------------------------------------

    if jenis == "pulang":

        if not presensi:

            return jsonify({
                "success": False,
                "message": (
                    "Anda belum melakukan "
                    "absen masuk."
                )
            })

        if not presensi["jam_masuk"]:

            return jsonify({
                "success": False,
                "message": (
                    "Anda belum melakukan "
                    "absen masuk."
                )
            })

        if presensi["jam_pulang"]:

            return jsonify({
                "success": False,
                "message": (
                    "Anda sudah melakukan "
                    "absen pulang hari ini."
                )
            })

    # ---------------------------------------------
    # SIMPAN FOTO
    # ---------------------------------------------

    try:

        if "," in foto:

            foto_data = foto.split(
                ",",
                1
            )[1]

        else:

            foto_data = foto

        image_data = base64.b64decode(
            foto_data
        )

        filename = (
            "presensi_"
            + str(perangkat_id)
            + "_"
            + jenis
            + "_"
            + uuid.uuid4().hex[:10]
            + ".jpg"
        )

        filepath = (
            UPLOAD_FOLDER
            / filename
        )

        with open(
            filepath,
            "wb"
        ) as file:

            file.write(
                image_data
            )

    except Exception as e:

        return jsonify({
            "success": False,
            "message": (
                "Gagal menyimpan foto: "
                + str(e)
            )
        })

    foto_path = (
        "uploads/presensi/"
        + filename
    )

    jam = jam_sekarang()

    # ---------------------------------------------
    # SIMPAN DATABASE
    # ---------------------------------------------

    conn = get_connection()

    try:

        if jenis == "masuk":

            if presensi is None:

                conn.execute("""
                    INSERT INTO presensi (
                        perangkat_id,
                        tanggal,
                        jam_masuk,
                        foto_masuk,
                        latitude_masuk,
                        longitude_masuk,
                        latitude,
                        longitude,
                        status
                    )

                    VALUES (
                        ?,
                        ?,
                        ?,
                        ?,
                        ?,
                        ?,
                        ?,
                        ?,
                        'Hadir'
                    )
                """, (
                    perangkat_id,
                    tanggal,
                    jam,
                    foto_path,
                    latitude,
                    longitude,
                    latitude,
                    longitude
                ))

            else:

                conn.execute("""
                    UPDATE presensi

                    SET
                        jam_masuk = ?,
                        foto_masuk = ?,
                        latitude_masuk = ?,
                        longitude_masuk = ?,
                        latitude = ?,
                        longitude = ?

                    WHERE id = ?

                """, (
                    jam,
                    foto_path,
                    latitude,
                    longitude,
                    latitude,
                    longitude,
                    presensi["id"]
                ))

        else:

            conn.execute("""
                UPDATE presensi

                SET
                    jam_pulang = ?,
                    foto_pulang = ?,
                    latitude_pulang = ?,
                    longitude_pulang = ?

                WHERE id = ?

            """, (
                jam,
                foto_path,
                latitude,
                longitude,
                presensi["id"]
            ))

        conn.commit()

    except Exception as e:

        conn.rollback()

        conn.close()

        try:

            if filepath.exists():

                filepath.unlink()

        except Exception:

            pass

        return jsonify({
            "success": False,
            "message": (
                "Gagal menyimpan presensi: "
                + str(e)
            )
        })

    conn.close()

    if jenis == "masuk":

        pesan = (
            "Absen masuk berhasil disimpan."
        )

    else:

        pesan = (
            "Absen pulang berhasil disimpan."
        )

    return jsonify({
        "success": True,
        "message": pesan,
        "jam": jam
    })


# =========================================================
# RIWAYAT PERANGKAT
# =========================================================

@app.route("/perangkat/riwayat")
def perangkat_riwayat():

    if "user_id" not in session:

        return redirect(
            url_for("login")
        )

    if session.get("role") != "perangkat":

        return redirect(
            url_for("admin_dashboard")
        )

    perangkat_id = session.get(
        "perangkat_id"
    )

    data = fetch_all("""
        SELECT *
        FROM presensi

        WHERE perangkat_id = ?

        ORDER BY
            tanggal DESC,
            id DESC

    """, (
        perangkat_id,
    ))

    return render_template(
        "perangkat_riwayat.html",
        data=data
    )


# =========================================================
# PROFIL PERANGKAT
# =========================================================

@app.route(
    "/perangkat/profil",
    methods=["GET", "POST"]
)
def perangkat_profil():

    if "user_id" not in session:

        return redirect(
            url_for("login")
        )

    if session.get("role") != "perangkat":

        return redirect(
            url_for("admin_dashboard")
        )

    perangkat_id = session.get(
        "perangkat_id"
    )

    conn = get_connection()

    perangkat = conn.execute("""
        SELECT
            perangkat.*,
            users.username

        FROM perangkat

        LEFT JOIN users
            ON users.perangkat_id =
               perangkat.id

        WHERE perangkat.id = ?

    """, (
        perangkat_id,
    )).fetchone()

    if perangkat is None:

        conn.close()

        session.clear()

        return redirect(
            url_for("login")
        )

    if request.method == "POST":

        foto = request.files.get(
            "foto"
        )

        if (
            foto is None
            or foto.filename == ""
        ):

            conn.close()

            return render_template(
                "perangkat_profil.html",
                perangkat=perangkat,
                error=(
                    "Silakan pilih foto terlebih dahulu."
                )
            )

        allowed_extensions = {
            "jpg",
            "jpeg",
            "png",
            "webp"
        }

        if "." not in foto.filename:

            conn.close()

            return render_template(
                "perangkat_profil.html",
                perangkat=perangkat,
                error=(
                    "Format foto tidak valid."
                )
            )

        extension = (
            foto.filename
            .rsplit(".", 1)[-1]
            .lower()
        )

        if extension not in allowed_extensions:

            conn.close()

            return render_template(
                "perangkat_profil.html",
                perangkat=perangkat,
                error=(
                    "Format foto harus "
                    "JPG, JPEG, PNG, atau WEBP."
                )
            )

        foto.seek(0, 2)

        ukuran_file = foto.tell()

        foto.seek(0)

        max_size = (
            5 * 1024 * 1024
        )

        if ukuran_file > max_size:

            conn.close()

            return render_template(
                "perangkat_profil.html",
                perangkat=perangkat,
                error=(
                    "Ukuran foto maksimal 5 MB."
                )
            )

        filename = (
            "profil_"
            + str(perangkat_id)
            + "_"
            + uuid.uuid4().hex[:10]
            + "."
            + extension
        )

        profile_folder = (
            BASE_DIR
            / "static"
            / "uploads"
            / "presensi"
        )

        profile_folder.mkdir(
            parents=True,
            exist_ok=True
        )

        filepath = (
            profile_folder
            / filename
        )

        try:

            foto.save(
                filepath
            )

        except Exception as e:

            conn.close()

            return render_template(
                "perangkat_profil.html",
                perangkat=perangkat,
                error=(
                    "Gagal menyimpan foto: "
                    + str(e)
                )
            )

        foto_path = (
            "uploads/presensi/"
            + filename
        )

        foto_lama = perangkat["foto"]

        try:

            conn.execute("""
                UPDATE perangkat

                SET foto = ?

                WHERE id = ?

            """, (
                foto_path,
                perangkat_id
            ))

            conn.commit()

        except Exception as e:

            conn.rollback()

            conn.close()

            try:

                if filepath.exists():

                    filepath.unlink()

            except Exception:

                pass

            return render_template(
                "perangkat_profil.html",
                perangkat=perangkat,
                error=(
                    "Gagal memperbarui database: "
                    + str(e)
                )
            )

        if foto_lama:

            try:

                old_filename = Path(
                    foto_lama
                ).name

                old_filepath = (
                    BASE_DIR
                    / "static"
                    / "uploads"
                    / "presensi"
                    / old_filename
                )

                if (
                    old_filepath.exists()
                    and old_filepath != filepath
                ):

                    old_filepath.unlink()

            except Exception:

                pass

        perangkat = conn.execute("""
            SELECT
                perangkat.*,
                users.username

            FROM perangkat

            LEFT JOIN users
                ON users.perangkat_id =
                   perangkat.id

            WHERE perangkat.id = ?

        """, (
            perangkat_id,
        )).fetchone()

        conn.close()

        return render_template(
            "perangkat_profil.html",
            perangkat=perangkat,
            success=(
                "Foto profil berhasil diperbarui."
            )
        )

    conn.close()

    return render_template(
        "perangkat_profil.html",
        perangkat=perangkat
    )


# =========================================================
# UBAH PASSWORD
# =========================================================

@app.route(
    "/perangkat/ubah-password",
    methods=["GET", "POST"]
)
def ubah_password():

    if "user_id" not in session:

        return redirect(
            url_for("login")
        )

    if session.get("role") != "perangkat":

        return redirect(
            url_for("admin_dashboard")
        )

    if request.method == "POST":

        password_lama = request.form.get(
            "password_lama",
            ""
        )

        password_baru = request.form.get(
            "password_baru",
            ""
        )

        password_konfirmasi = request.form.get(
            "password_konfirmasi",
            ""
        )

        user = fetch_one("""
            SELECT *
            FROM users
            WHERE id = ?
        """, (
            session["user_id"],
        ))

        if user is None:

            return redirect(
                url_for("logout")
            )

        if not check_password_hash(
            user["password"],
            password_lama
        ):

            return render_template(
                "ubah_password.html",
                error=(
                    "Password lama salah."
                )
            )

        if len(password_baru) < 6:

            return render_template(
                "ubah_password.html",
                error=(
                    "Password baru minimal "
                    "6 karakter."
                )
            )

        if (
            password_baru
            != password_konfirmasi
        ):

            return render_template(
                "ubah_password.html",
                error=(
                    "Konfirmasi password "
                    "tidak sama."
                )
            )

        password_hash = (
            generate_password_hash(
                password_baru
            )
        )

        execute("""
            UPDATE users

            SET password = ?

            WHERE id = ?

        """, (
            password_hash,
            session["user_id"]
        ))

        return render_template(
            "ubah_password.html",
            success=(
                "Password berhasil diubah."
            )
        )

    return render_template(
        "ubah_password.html"
    )


# =========================================================
# JALANKAN APLIKASI
# =========================================================

if __name__ == "__main__":

    app.run(
        debug=True,
        host="127.0.0.1",
        port=5000
    )